import os
from zipfile import ZipFile
from xml.etree import ElementTree


IMAGE_TYPES = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg"}
MAX_VISUAL_PAGES = 40


def extract_planner_material(path, analyze_image):
    """Extract text and visible diagram evidence without dropping scanned pages."""
    extension = os.path.splitext(path)[1].lower()
    sections, visuals, warnings = [], [], []
    image_count = 0

    def record_relationship(description, location):
        visuals.append({"kind": "relationship", "description": description, "location": location})
        sections.append(f"{location} - relationship: {description}")

    def read_smartart(archive):
        namespaces = {"d": "http://schemas.openxmlformats.org/drawingml/2006/diagram", "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        for filename in archive.namelist():
            if "/diagrams/data" not in filename or not filename.endswith(".xml"):
                continue
            root = ElementTree.fromstring(archive.read(filename))
            labels = {point.get("modelId"): " ".join(point.itertext()).strip() for point in root.findall(".//d:pt", namespaces)}
            for connection in root.findall(".//d:cxn", namespaces):
                left, right = labels.get(connection.get("srcId")), labels.get(connection.get("destId"))
                if left and right:
                    record_relationship(f"Diagram connects {left} with {right}.", os.path.basename(filename))

    def read_image(content, mime, location):
        nonlocal image_count
        image_count += 1
        if image_count > MAX_VISUAL_PAGES:
            raise ValueError("This file has more than 40 visual pages or images. Split it into smaller files.")
        result = analyze_image(content, mime)
        text = str(result.get("text") or "").strip()
        diagrams = result.get("diagrams", [])
        relationships = result.get("relationships", [])
        if text:
            sections.append(f"{location} - visible text:\n{text}")
        for category, values in [("diagram", diagrams), ("relationship", relationships)]:
            if not isinstance(values, list):
                continue
            for value in values:
                if isinstance(value, str) and value.strip():
                    visuals.append({"kind": category, "description": value.strip(), "location": location})
                    sections.append(f"{location} - {category}: {value.strip()}")

    if extension in IMAGE_TYPES:
        import fitz
        with fitz.open(path) as image:
            page = image[0]
            pixmap = page.get_pixmap(matrix=fitz.Matrix(min(2, 1800 / max(page.rect.width, page.rect.height)), min(2, 1800 / max(page.rect.width, page.rect.height))), alpha=False)
            read_image(pixmap.tobytes("png"), "image/png", "Image")
    elif extension == ".pdf":
        import fitz
        with fitz.open(path) as document:
            if document.needs_pass:
                raise ValueError("Unlock this PDF before uploading it.")
            for index, page in enumerate(document):
                text = page.get_text("text").strip()
                location = f"Page {index + 1}"
                if text:
                    sections.append(f"{location}:\n{text}")
                if not text or page.get_images() or page.get_drawings():
                    scale = min(2, 1800 / max(page.rect.width, page.rect.height))
                    read_image(page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False).tobytes("png"), "image/png", location)
    elif extension == ".docx":
        from docx import Document
        document = Document(path)
        sections.extend(p.text for p in document.paragraphs if p.text.strip())
        for table in document.tables:
            sections.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        with ZipFile(path) as archive:
            read_smartart(archive)
            for name in archive.namelist():
                if not name.startswith("word/media/"):
                    continue
                mime = IMAGE_TYPES.get(os.path.splitext(name)[1].lower())
                if mime:
                    read_image(archive.read(name), mime, f"Embedded image {os.path.basename(name)}")
                else:
                    warnings.append(f"Embedded image {os.path.basename(name)} could not be read. Upload it as PNG or JPEG.")
    elif extension == ".pptx":
        from pptx import Presentation
        def read_shapes(shapes, location):
            labels = {shape.shape_id: shape.text.strip() for shape in shapes if shape.has_text_frame and shape.text.strip()}
            for shape in shapes:
                if shape.has_text_frame and shape.text.strip():
                    sections.append(f"{location}: {shape.text}")
                if shape.has_table:
                    sections.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
                if shape.has_chart:
                    chart = shape.chart
                    for series in chart.series:
                        sections.append(f"{location} chart - {series.name}: {list(series.values)}")
                    try:
                        sections.append(f"{location} chart categories: {', '.join(str(c.label) for c in chart.plots[0].categories)}")
                    except (AttributeError, IndexError):
                        pass
                starts = shape.element.xpath(".//a:stCxn")
                ends = shape.element.xpath(".//a:endCxn")
                if starts and ends:
                    left, right = labels.get(int(starts[0].get("id", 0))), labels.get(int(ends[0].get("id", 0)))
                    if left and right:
                        record_relationship(f"A connector joins {left} with {right}.", location)
                if hasattr(shape, "image"):
                    mime = shape.image.content_type
                    if mime in IMAGE_TYPES.values():
                        read_image(shape.image.blob, mime, location)
                    else:
                        warnings.append(f"{location}: an image could not be read. Upload it as PNG or JPEG.")
                if hasattr(shape, "shapes"):
                    read_shapes(shape.shapes, location)
        for index, slide in enumerate(Presentation(path).slides):
            read_shapes(slide.shapes, f"Slide {index + 1}")
        with ZipFile(path) as archive:
            read_smartart(archive)
    elif extension == ".txt":
        with open(path, encoding="utf-8", errors="replace") as source:
            sections.append(source.read())
    else:
        raise ValueError("Unsupported study material type.")

    text = "\n\n".join(sections).strip()
    if not text:
        raise ValueError("No readable text or diagram information was found. Upload a clearer image or document.")
    return {"text": text, "visualFindings": visuals, "warnings": warnings}
